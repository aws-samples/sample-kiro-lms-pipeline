"""Tests for scripts/build_assessment.py (validate + XLSX export).

The script's optional dependencies (pyyaml, openpyxl) are imported lazily,
so the module itself always imports. Test classes that exercise YAML loading
or XLSX writing are skipped when the corresponding package is missing.
"""

from __future__ import annotations

import importlib.util
from pathlib import Path

import pytest

from build_assessment import (
    Assessment,
    ModuleInfo,
    Option,
    Question,
    Section,
    discover_modules,
    extract_knowledge_checks,
    format_correct_answer,
    parse_args,
    parse_assessment_yaml,
    resolve_config,
    strip_frontmatter,
    validate_assessment,
    write_metadata_xlsx,
    write_xlsx,
)

FIXTURES = Path(__file__).parent / "fixtures"

requires_yaml = pytest.mark.skipif(
    importlib.util.find_spec("yaml") is None, reason="pyyaml not installed"
)
requires_openpyxl = pytest.mark.skipif(
    importlib.util.find_spec("openpyxl") is None, reason="openpyxl not installed"
)


def _question(response_type: str, options: list[Option], correct_answer: str = "", **kwargs) -> Question:
    return Question(
        question_text="Sample question?",
        response_type=response_type,
        options=options,
        correct_answer=correct_answer,
        **kwargs,
    )


# ---------------------------------------------------------------------------
# parse_args
# ---------------------------------------------------------------------------


class TestParseArgsDefaults:
    def test_default_course_name(self):
        args = parse_args([])
        assert args.course_name == "sample-course"

    def test_metadata_defaults_are_none(self):
        args = parse_args([])
        assert args.name is None
        assert args.description is None
        assert args.passing_score is None
        assert args.time_limit is None
        assert args.tags is None
        assert args.language is None


class TestParseArgsExplicit:
    def test_positional_course_name(self):
        args = parse_args(["my-course"])
        assert args.course_name == "my-course"

    def test_name_and_description(self):
        args = parse_args(["--name", "Final Exam", "--description", "End-of-course check"])
        assert args.name == "Final Exam"
        assert args.description == "End-of-course check"

    def test_passing_score_is_int(self):
        args = parse_args(["--passing-score", "85"])
        assert args.passing_score == 85

    def test_time_limit_is_int(self):
        args = parse_args(["--time-limit", "45"])
        assert args.time_limit == 45

    def test_tags_single(self):
        args = parse_args(["--tags", "scorm"])
        assert args.tags == ["scorm"]

    def test_tags_multiple(self):
        args = parse_args(["--tags", "scorm", "authoring", "sample"])
        assert args.tags == ["scorm", "authoring", "sample"]

    def test_language(self):
        args = parse_args(["--language", "Spanish (ES)"])
        assert args.language == "Spanish (ES)"

    def test_full_invocation(self):
        args = parse_args(
            [
                "my-course",
                "--name", "Final Exam",
                "--description", "End-of-course check",
                "--passing-score", "90",
                "--time-limit", "60",
                "--tags", "a", "b",
                "--language", "English (US)",
            ]
        )
        assert args.course_name == "my-course"
        assert args.name == "Final Exam"
        assert args.description == "End-of-course check"
        assert args.passing_score == 90
        assert args.time_limit == 60
        assert args.tags == ["a", "b"]
        assert args.language == "English (US)"


# ---------------------------------------------------------------------------
# resolve_config
# ---------------------------------------------------------------------------


class TestResolveConfig:
    def test_defaults_with_empty_config(self):
        assessment = resolve_config(parse_args([]), {})
        assert assessment.name == ""
        assert assessment.description == ""
        assert assessment.passing_score == 80
        assert assessment.time_limit_minutes is None
        assert assessment.language == "English (US)"
        assert assessment.tags == []

    def test_course_yaml_assessment_block(self):
        course_config = {
            "assessment": {
                "name": "Course Check",
                "description": "From course.yaml",
                "passing_score": 70,
                "time_limit_minutes": 15,
                "language": "Spanish (ES)",
                "tags": ["from-course"],
            }
        }
        assessment = resolve_config(parse_args([]), course_config)
        assert assessment.name == "Course Check"
        assert assessment.description == "From course.yaml"
        assert assessment.passing_score == 70
        assert assessment.time_limit_minutes == 15
        assert assessment.language == "Spanish (ES)"
        assert assessment.tags == ["from-course"]

    def test_course_level_fallbacks(self):
        course_config = {
            "course": {"title": "Fallback Title", "tags": ["course-tag"]},
            "scorm": {"passing_score": 100},
        }
        assessment = resolve_config(parse_args([]), course_config)
        assert assessment.name == "Fallback Title"
        assert assessment.passing_score == 100
        assert assessment.tags == ["course-tag"]

    def test_yaml_metadata_overrides_course_yaml(self):
        course_config = {"assessment": {"name": "Course Check", "passing_score": 70}}
        yaml_metadata = {"name": "Authored Check", "passing_score": 85}
        assessment = resolve_config(parse_args([]), course_config, yaml_metadata)
        assert assessment.name == "Authored Check"
        assert assessment.passing_score == 85

    def test_cli_overrides_everything(self):
        course_config = {"assessment": {"name": "Course Check", "passing_score": 70}}
        yaml_metadata = {"name": "Authored Check", "passing_score": 85}
        args = parse_args(["--name", "CLI Check", "--passing-score", "95"])
        assessment = resolve_config(args, course_config, yaml_metadata)
        assert assessment.name == "CLI Check"
        assert assessment.passing_score == 95


# ---------------------------------------------------------------------------
# discover_modules
# ---------------------------------------------------------------------------


class TestDiscoverModules:
    def _make_course(self, tmp_path: Path, module_ids: list[str]) -> Path:
        course_dir = tmp_path / "course"
        for module_id in module_ids:
            lesson = course_dir / "build" / module_id / "lesson.md"
            lesson.parent.mkdir(parents=True, exist_ok=True)
            lesson.write_text(f"# Lesson for {module_id}\n")
        return course_dir

    def test_discovers_modules_with_lessons(self, tmp_path):
        course_dir = self._make_course(tmp_path, ["mod-1", "mod-2"])
        config = {
            "modules": [
                {"id": "mod-1", "title": "Module One"},
                {"id": "mod-2", "title": "Module Two"},
            ]
        }
        modules = discover_modules(course_dir, config)
        assert modules == [
            ModuleInfo("mod-1", "Module One", course_dir / "build" / "mod-1" / "lesson.md"),
            ModuleInfo("mod-2", "Module Two", course_dir / "build" / "mod-2" / "lesson.md"),
        ]

    def test_skips_resources_only_modules(self, tmp_path, capsys):
        course_dir = self._make_course(tmp_path, ["mod-1", "mod-2"])
        config = {
            "modules": [
                {"id": "mod-1", "title": "Module One"},
                {"id": "mod-2", "title": "Resources", "resources_only": True},
            ]
        }
        modules = discover_modules(course_dir, config)
        assert [m.id for m in modules] == ["mod-1"]
        assert "resources_only" in capsys.readouterr().err

    def test_skips_missing_lesson_files(self, tmp_path, capsys):
        course_dir = self._make_course(tmp_path, ["mod-2"])
        config = {
            "modules": [
                {"id": "mod-1", "title": "Module One"},
                {"id": "mod-2", "title": "Module Two"},
            ]
        }
        modules = discover_modules(course_dir, config)
        assert [m.id for m in modules] == ["mod-2"]
        err = capsys.readouterr().err
        assert "WARNING" in err
        assert "mod-1" in err

    def test_empty_for_no_modules(self, tmp_path):
        assert discover_modules(tmp_path, {"modules": []}) == []

    def test_empty_for_missing_modules_key(self, tmp_path):
        assert discover_modules(tmp_path, {}) == []

    def test_skips_modules_without_id_or_title(self, tmp_path):
        course_dir = self._make_course(tmp_path, ["mod-1"])
        config = {
            "modules": [
                {"title": "No id"},
                {"id": "no-title"},
                {"id": "mod-1", "title": "Module One"},
            ]
        }
        modules = discover_modules(course_dir, config)
        assert [m.id for m in modules] == ["mod-1"]


# ---------------------------------------------------------------------------
# strip_frontmatter
# ---------------------------------------------------------------------------


class TestStripFrontmatter:
    def test_strips_frontmatter(self):
        content = "---\ntitle: Test\n---\n# Body\n"
        assert strip_frontmatter(content) == "# Body\n"

    def test_no_frontmatter_returned_as_is(self):
        content = "# Body\nNo frontmatter here.\n"
        assert strip_frontmatter(content) == content

    def test_unclosed_frontmatter_returned_as_is(self):
        content = "---\ntitle: Test\n# Body without closing delimiter\n"
        assert strip_frontmatter(content) == content

    def test_bare_dashes_only(self):
        assert strip_frontmatter("---") == "---"

    def test_multiline_frontmatter(self):
        content = "---\ntitle: Test\ntags:\n  - a\n  - b\n---\nBody text\n"
        assert strip_frontmatter(content) == "Body text\n"


# ---------------------------------------------------------------------------
# validate_assessment
# ---------------------------------------------------------------------------


def _valid_assessment(**overrides) -> Assessment:
    q = _question(
        "Single select",
        [Option(text="A", correct=True), Option(text="B", correct=False)],
    )
    fields = dict(
        name="Valid Assessment",
        description="A valid assessment.",
        passing_score=80,
        sections=[Section(title="Section 1", questions=[q])],
    )
    fields.update(overrides)
    return Assessment(**fields)


class TestValidateAssessment:
    def test_valid_assessment_has_no_errors(self):
        assert validate_assessment(_valid_assessment()) == []

    def test_name_too_long(self):
        errors = validate_assessment(_valid_assessment(name="x" * 251))
        assert any(e.field == "name" and "250" in e.message for e in errors)

    def test_description_too_long(self):
        errors = validate_assessment(_valid_assessment(description="x" * 5001))
        assert any(e.field == "description" and "5000" in e.message for e in errors)

    def test_too_many_tags(self):
        errors = validate_assessment(_valid_assessment(tags=[f"t{i}" for i in range(16)]))
        assert any(e.field == "tags" for e in errors)

    def test_passing_score_out_of_range(self):
        assert any(e.field == "passing_score" for e in validate_assessment(_valid_assessment(passing_score=101)))
        assert any(e.field == "passing_score" for e in validate_assessment(_valid_assessment(passing_score=-1)))

    def test_zero_questions(self):
        errors = validate_assessment(_valid_assessment(sections=[Section(title="Empty")]))
        assert any(e.field == "questions" and "zero" in e.message for e in errors)

    def test_too_many_questions_total(self):
        sections = [
            Section(title=f"S{i}", questions=[
                Question(question_text=f"Q{i}-{j}", response_type="Free text", options=[], correct_answer="")
                for j in range(100)
            ])
            for i in range(8)
        ]
        errors = validate_assessment(_valid_assessment(sections=sections))
        assert any(e.field == "questions" and "750" in e.message for e in errors)

    def test_section_question_cap(self):
        questions = [
            Question(question_text=f"Q{j}", response_type="Free text", options=[], correct_answer="")
            for j in range(101)
        ]
        errors = validate_assessment(_valid_assessment(sections=[Section(title="Big", questions=questions)]))
        assert any("100 questions" in e.message for e in errors)

    def test_question_text_too_long(self):
        q = Question(question_text="x" * 2001, response_type="Free text", options=[], correct_answer="")
        errors = validate_assessment(_valid_assessment(sections=[Section(title="S", questions=[q])]))
        assert any("Question text exceeds 2000" in e.message for e in errors)

    def test_option_text_too_long(self):
        q = _question(
            "Single select",
            [Option(text="x" * 2001, correct=True), Option(text="B", correct=False)],
        )
        errors = validate_assessment(_valid_assessment(sections=[Section(title="S", questions=[q])]))
        assert any("Option text exceeds 2000" in e.message for e in errors)

    def test_too_many_options(self):
        q = _question(
            "Single select",
            [Option(text=f"O{i}", correct=(i == 0)) for i in range(21)],
        )
        errors = validate_assessment(_valid_assessment(sections=[Section(title="S", questions=[q])]))
        assert any("more than 20 options" in e.message for e in errors)

    def test_select_types_need_two_options(self):
        for rt in ("Single select", "Multi select"):
            q = _question(rt, [Option(text="Only", correct=True)])
            errors = validate_assessment(_valid_assessment(sections=[Section(title="S", questions=[q])]))
            assert any("at least 2 options" in e.message for e in errors), rt

    def test_reports_all_violations_at_once(self):
        q = Question(question_text="x" * 2001, response_type="Single select",
                     options=[Option(text="Only", correct=True)], correct_answer="")
        errors = validate_assessment(
            _valid_assessment(name="x" * 251, passing_score=150,
                              sections=[Section(title="S", questions=[q])])
        )
        fields = {e.field for e in errors}
        assert "name" in fields
        assert "passing_score" in fields
        assert len(errors) >= 4


# ---------------------------------------------------------------------------
# format_correct_answer
# ---------------------------------------------------------------------------


class TestFormatCorrectAnswer:
    def test_single_select_returns_1_based_int_index(self):
        q = _question(
            "Single select",
            [Option(text="A", correct=False), Option(text="B", correct=True), Option(text="C", correct=False)],
        )
        assert format_correct_answer(q) == 2

    def test_single_select_first_option(self):
        q = _question("Single select", [Option(text="A", correct=True), Option(text="B", correct=False)])
        assert format_correct_answer(q) == 1

    def test_single_select_fallback_to_correct_answer(self):
        q = _question("Single select", [Option(text="A", correct=False), Option(text="B", correct=False)],
                      correct_answer="3")
        assert format_correct_answer(q) == 3

    def test_multi_select_json_list_string(self):
        q = _question(
            "Multi select",
            [
                Option(text="A", correct=True),
                Option(text="B", correct=False),
                Option(text="C", correct=True),
                Option(text="D", correct=True),
            ],
        )
        assert format_correct_answer(q) == "[1, 3, 4]"

    def test_text_json_list_of_strings(self):
        q = _question("Text", [Option(text="Python", correct=True), Option(text="Java", correct=True)])
        assert format_correct_answer(q) == '["Python", "Java"]'

    def test_yes_no_from_options(self):
        q = _question("Yes/No", [Option(text="Yes", correct=True), Option(text="No", correct=False)])
        assert format_correct_answer(q) == "YES"

    def test_yes_no_from_correct_answer(self):
        q = _question("Yes/No", [], correct_answer="no")
        assert format_correct_answer(q) == "NO"

    def test_true_false_from_options(self):
        q = _question("True/False", [Option(text="True", correct=False), Option(text="False", correct=True)])
        assert format_correct_answer(q) == "FALSE"

    def test_true_false_from_correct_answer(self):
        q = _question("True/False", [], correct_answer="true")
        assert format_correct_answer(q) == "TRUE"

    def test_pass_fail_from_options(self):
        q = _question("Pass/Fail", [Option(text="Pass", correct=True), Option(text="Fail", correct=False)])
        assert format_correct_answer(q) == "PASS"

    def test_number_and_percentage_as_is(self):
        assert format_correct_answer(_question("Number", [], correct_answer="30")) == "30"
        assert format_correct_answer(_question("Percentage", [], correct_answer="50")) == "50"

    def test_free_text_as_is(self):
        assert format_correct_answer(_question("Free text", [], correct_answer="")) == ""


# ---------------------------------------------------------------------------
# parse_assessment_yaml
# ---------------------------------------------------------------------------


@requires_yaml
class TestParseAssessmentYaml:
    def test_parses_fixture(self):
        sections, metadata = parse_assessment_yaml(FIXTURES / "sample_assessment.yaml")

        assert metadata["name"] == "SCORM Authoring Fundamentals Check"
        assert metadata["passing_score"] == 80
        assert metadata["time_limit_minutes"] == 30
        assert metadata["language"] == "English (US)"

        # 3 declared sections + "General" created for the unsectioned question
        titles = [s.title for s in sections]
        assert titles == [
            "The Authoring Pipeline",
            "Quiz Interactions",
            "Media and Completion Rules",
            "General",
        ]
        assert [len(s.questions) for s in sections] == [2, 1, 1, 1]

        q1 = sections[0].questions[0]
        assert q1.response_type == "Single select"
        assert q1.options[0].correct is True
        assert q1.score == 1

    def test_question_with_unknown_section_creates_it(self, tmp_path):
        path = tmp_path / "assessment.yaml"
        path.write_text(
            "questions:\n"
            "  - question_text: 'Q?'\n"
            "    response_type: 'Free text'\n"
            "    section: 'Brand New'\n"
        )
        sections, _ = parse_assessment_yaml(path)
        assert [s.title for s in sections] == ["Brand New"]
        assert len(sections[0].questions) == 1

    def test_question_without_section_goes_to_general(self, tmp_path):
        path = tmp_path / "assessment.yaml"
        path.write_text(
            "questions:\n"
            "  - question_text: 'Q?'\n"
            "    response_type: 'Free text'\n"
        )
        sections, _ = parse_assessment_yaml(path)
        assert [s.title for s in sections] == ["General"]

    def test_invalid_response_type_exits(self, tmp_path, capsys):
        path = tmp_path / "assessment.yaml"
        path.write_text(
            "questions:\n"
            "  - question_text: 'Q?'\n"
            "    response_type: 'Essay'\n"
        )
        with pytest.raises(SystemExit):
            parse_assessment_yaml(path)
        err = capsys.readouterr().err
        assert "invalid response_type" in err
        assert "Essay" in err

    def test_missing_question_text_exits(self, tmp_path, capsys):
        path = tmp_path / "assessment.yaml"
        path.write_text("questions:\n  - response_type: 'Free text'\n")
        with pytest.raises(SystemExit):
            parse_assessment_yaml(path)
        assert "question_text" in capsys.readouterr().err

    def test_malformed_yaml_exits(self, tmp_path, capsys):
        path = tmp_path / "assessment.yaml"
        path.write_text("questions:\n  - question_text: 'unclosed\n")
        with pytest.raises(SystemExit):
            parse_assessment_yaml(path)
        assert "ERROR" in capsys.readouterr().err

    def test_missing_file_exits(self, tmp_path):
        with pytest.raises(SystemExit):
            parse_assessment_yaml(tmp_path / "nope.yaml")


# ---------------------------------------------------------------------------
# extract_knowledge_checks
# ---------------------------------------------------------------------------


@requires_yaml
class TestExtractKnowledgeChecks:
    def test_extracts_from_fixture_lesson(self):
        content = (FIXTURES / "sample_lesson.md").read_text()
        questions = extract_knowledge_checks(content, module_id="module-1")

        assert len(questions) == 2

        q1 = questions[0]
        assert "source of truth" in q1.question_text
        assert q1.response_type == "Single select"
        assert len(q1.options) == 3
        assert q1.correct_answer == "1"
        assert q1.options[0].correct is True

        q2 = questions[1]
        assert "heading" in q2.question_text
        assert q2.response_type == "Single select"
        assert len(q2.options) == 4
        assert q2.correct_answer == "2"

    def test_multi_select_when_multiple_correct(self):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'Pick all that apply'\n"
            "options:\n"
            "  - text: 'A'\n"
            "    correct: true\n"
            "  - text: 'B'\n"
            "    correct: true\n"
            "  - text: 'C'\n"
            "    correct: false\n"
            "{% endknowledgeCheck %}\n"
        )
        questions = extract_knowledge_checks(content)
        assert len(questions) == 1
        assert questions[0].response_type == "Multi select"
        assert questions[0].correct_answer == "[1, 2]"

    def test_feedback_becomes_option_explanation(self):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'Q?'\n"
            "options:\n"
            "  - text: 'A'\n"
            "    correct: true\n"
            "    feedback: 'Right you are.'\n"
            "  - text: 'B'\n"
            "    correct: false\n"
            "    feedback: 'Not quite.'\n"
            "{% endknowledgeCheck %}\n"
        )
        questions = extract_knowledge_checks(content)
        assert questions[0].options[0].explanation == "Right you are."
        assert questions[0].options[1].explanation == "Not quite."

    def test_correct_answer_explanation_joins_correct_feedback(self):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'Q?'\n"
            "options:\n"
            "  - text: 'A'\n"
            "    correct: true\n"
            "    feedback: 'First reason.'\n"
            "  - text: 'B'\n"
            "    correct: true\n"
            "    feedback: 'Second reason.'\n"
            "  - text: 'C'\n"
            "    correct: false\n"
            "    feedback: 'Wrong.'\n"
            "{% endknowledgeCheck %}\n"
        )
        questions = extract_knowledge_checks(content)
        assert questions[0].correct_answer_explanation == "First reason. Second reason."

    def test_special_yaml_characters(self):
        content = (
            "{% knowledgeCheck %}\n"
            'question: "What does `--dry-run: preview` do?"\n'
            "options:\n"
            '  - text: "Previews changes: nothing is written"\n'
            "    correct: true\n"
            '  - text: "Writes everything"\n'
            "    correct: false\n"
            "{% endknowledgeCheck %}\n"
        )
        questions = extract_knowledge_checks(content)
        assert len(questions) == 1
        assert "preview" in questions[0].question_text

    def test_malformed_yaml_warns_and_continues(self, capsys):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'broken\n"
            "  bad: [indentation\n"
            "{% endknowledgeCheck %}\n"
            "{% knowledgeCheck %}\n"
            "question: 'Good one?'\n"
            "options:\n"
            "  - text: 'A'\n"
            "    correct: true\n"
            "  - text: 'B'\n"
            "    correct: false\n"
            "{% endknowledgeCheck %}\n"
        )
        questions = extract_knowledge_checks(content, module_id="mod-x")
        assert len(questions) == 1
        err = capsys.readouterr().err
        assert "WARNING" in err
        assert "[mod-x]" in err

    def test_missing_question_field_warns(self, capsys):
        content = (
            "{% knowledgeCheck %}\n"
            "options:\n"
            "  - text: 'A'\n"
            "    correct: true\n"
            "{% endknowledgeCheck %}\n"
        )
        assert extract_knowledge_checks(content) == []
        assert "missing 'question'" in capsys.readouterr().err

    def test_missing_options_warns(self, capsys):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'Q?'\n"
            "{% endknowledgeCheck %}\n"
        )
        assert extract_knowledge_checks(content) == []
        assert "options" in capsys.readouterr().err

    def test_unclosed_block_warns(self, capsys):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'Q?'\n"
            "options:\n"
            "  - text: 'A'\n"
            "    correct: true\n"
        )
        assert extract_knowledge_checks(content) == []
        assert "no closing tag" in capsys.readouterr().err

    def test_fewer_than_two_options_warns(self, capsys):
        content = (
            "{% knowledgeCheck %}\n"
            "question: 'Q?'\n"
            "options:\n"
            "  - text: 'Only one'\n"
            "    correct: true\n"
            "{% endknowledgeCheck %}\n"
        )
        assert extract_knowledge_checks(content) == []
        assert "fewer than 2 options" in capsys.readouterr().err

    def test_no_blocks_returns_empty_list(self):
        assert extract_knowledge_checks("# Just a lesson\n\nNo checks here.\n") == []


# ---------------------------------------------------------------------------
# XLSX export
# ---------------------------------------------------------------------------


def _xlsx_assessment(language: str = "English (US)") -> Assessment:
    single = _question(
        "Single select",
        [Option(text="A", correct=False), Option(text="B", correct=True), Option(text="C", correct=False)],
        correct_answer_explanation="B is right.",
        score=2,
    )
    multi = _question(
        "Multi select",
        [
            Option(text="W", correct=False),
            Option(text="X", correct=True),
            Option(text="Y", correct=False),
            Option(text="Z", correct=True),
        ],
    )
    return Assessment(
        name="XLSX Test Assessment",
        description="For XLSX writer tests.",
        language=language,
        sections=[
            Section(title="Section One", description="First section", questions=[single]),
            Section(title="Section Two", description="Second section", questions=[multi]),
        ],
    )


@requires_openpyxl
class TestWriteXlsx:
    def _load(self, path: Path):
        import openpyxl

        return openpyxl.load_workbook(str(path))

    def test_creates_file(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        assert out.exists()

    def test_creates_nested_output_dirs(self, tmp_path):
        out = tmp_path / "deep" / "nested" / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        assert out.exists()

    def test_sheet_names(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        wb = self._load(out)
        assert wb.sheetnames == ["Instructions", "Sections", "English (US) (d)"]

    def test_custom_language_sheet_name(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(language="Spanish (ES)"), out)
        wb = self._load(out)
        assert "Spanish (ES) (d)" in wb.sheetnames

    def test_instructions_sheet_content(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        ws = self._load(out)["Instructions"]
        assert "Excel file" in ws.cell(row=1, column=1).value

    def test_sections_sheet_rows(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        ws = self._load(out)["Sections"]
        header = [c.value for c in ws[1]]
        assert header == [
            "Language",
            "Section 1 title", "Section 1 description",
            "Section 2 title", "Section 2 description",
        ]
        data = [c.value for c in ws[2]]
        assert data == [
            "English (US)",
            "Section One", "First section",
            "Section Two", "Second section",
        ]

    def test_question_sheet_header_has_47_columns(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        ws = self._load(out)["English (US) (d)"]
        header = [c.value for c in ws[1]]
        assert len(header) == 47
        assert header[:4] == ["Section no", "Question id", "Question text", "Response type"]
        assert header[4] == "Option 1"
        assert header[5] == "Option 1 explanation"
        assert header[-3:] == ["Correct answer", "Correct answer explanation", "Question score"]

    def test_single_select_row(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        ws = self._load(out)["English (US) (d)"]
        row = [c.value for c in ws[2]]
        assert row[0] == 1  # Section no
        assert row[1] is None  # Question id blank for new questions
        assert row[3] == "Single select"
        assert row[-3] == 2  # numeric 1-based index of the correct option
        assert row[-2] == "B is right."
        assert row[-1] == 2  # score

    def test_multi_select_row_in_section_two(self, tmp_path):
        out = tmp_path / "assessment.xlsx"
        write_xlsx(_xlsx_assessment(), out)
        ws = self._load(out)["English (US) (d)"]
        row = [c.value for c in ws[3]]
        assert row[0] == 2  # Section no
        assert row[-3] == "[2, 4]"

    def test_metadata_xlsx(self, tmp_path):
        out = tmp_path / "metadata.xlsx"
        write_metadata_xlsx(_xlsx_assessment(), out)
        wb = self._load(out)
        assert wb.sheetnames == ["Instructions", "Metadata"]
        ws = wb["Metadata"]
        assert [c.value for c in ws[1]] == ["Language", "Title", "Description"]
        assert [c.value for c in ws[2]] == [
            "English (US)", "XLSX Test Assessment", "For XLSX writer tests.",
        ]


# ---------------------------------------------------------------------------
# End-to-end: fixture yaml -> validated -> XLSX
# ---------------------------------------------------------------------------


@requires_yaml
@requires_openpyxl
class TestEndToEnd:
    def test_fixture_yaml_validates_and_exports(self, tmp_path):
        yaml = pytest.importorskip("yaml")

        sections, metadata = parse_assessment_yaml(FIXTURES / "sample_assessment.yaml")
        course_config = yaml.safe_load((FIXTURES / "sample_course.yaml").read_text())
        assessment = resolve_config(parse_args([]), course_config, metadata)
        assessment.sections = sections

        assert validate_assessment(assessment) == []

        out = tmp_path / "assessment.xlsx"
        write_xlsx(assessment, out)
        assert out.exists()

        import openpyxl

        wb = openpyxl.load_workbook(str(out))
        ws = wb["English (US) (d)"]
        # 1 header row + 5 question rows
        assert ws.max_row == 6
