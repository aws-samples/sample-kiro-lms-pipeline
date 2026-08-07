#!/usr/bin/env python3
"""Extract speaker notes from a .pptx deck, keyed by slide number.

The notes embedded in a deck are usually the closest thing to authored
prose the deck carries - the backbone for turning slides into lesson
text (see the power's ``source-intake.md`` steering file). This script
emits each slide's notes as markdown (default) or plain text, numbered
densely so notes line up with the PNGs produced by
``scripts/extract_visuals.py`` (``slide-001.png`` <-> ``## Slide 1``).

Hidden slides are skipped by default. ``extract_visuals.py`` renders via
LibreOffice, whose PDF export omits hidden slides, so its PNGs cover only
the visible slides. To keep the note numbers aligned with those PNGs, this
script skips hidden slides too and renumbers the survivors from 1 - and
prints how many it skipped so the divergence is never silent. Pass
``--include-hidden`` to keep every slide (note: the numbers will then no
longer match the visuals for a deck that has hidden slides).

Requirements (optional-tier, not needed for the SCORM pipeline itself):

- python-pptx: ``pip install python-pptx==1.0.2``

Only ``.pptx`` is supported; export Keynote/Google Slides decks to
``.pptx`` first.

Usage:
    python3 scripts/extract_notes.py <deck.pptx> [--format markdown|text] [--include-hidden] [-o <file>]
"""

from __future__ import annotations

import argparse
import sys
from pathlib import Path


def fail(msg: str):
    print(f"ERROR: {msg}", file=sys.stderr)
    raise SystemExit(1)


def is_hidden(slide) -> bool:
    """True when a slide is marked "Hide Slide" in the deck.

    The state lives on the ``show`` attribute of the ``<p:sld>`` element:
    ``show="0"`` means hidden; an absent attribute means visible (the default).
    python-pptx has no public accessor for it, so read the underlying element.
    """
    return slide._element.get("show") == "0"


def extract_notes(deck: Path, include_hidden: bool = False) -> tuple[list[tuple[int, str]], int]:
    """Return ``(notes, hidden_skipped)``.

    ``notes`` is a list of ``(slide_number, notes_text)`` for every slide that
    will be emitted, numbered densely from 1 so the numbers line up with the
    PNGs from ``extract_visuals.py``. Slides without notes yield an empty string.

    By default hidden slides are skipped (``include_hidden=False``) to match
    LibreOffice's PDF export, which omits them - this keeps the note numbers
    aligned with the slide PNGs. ``hidden_skipped`` is the count that were
    dropped, so the caller can report the divergence instead of hiding it.
    Pass ``include_hidden=True`` to keep every slide.
    """
    try:
        from pptx import Presentation
    except ImportError:
        fail(
            "python-pptx not installed (needed to read .pptx speaker notes).\n"
            "Install with: pip install python-pptx==1.0.2"
        )
    from pptx import Presentation

    prs = Presentation(str(deck))
    notes: list[tuple[int, str]] = []
    hidden_skipped = 0
    number = 0
    for slide in prs.slides:
        if is_hidden(slide) and not include_hidden:
            hidden_skipped += 1
            continue
        number += 1
        text = ""
        if slide.has_notes_slide:
            frame = slide.notes_slide.notes_text_frame
            if frame is not None:
                text = frame.text.strip()
        notes.append((number, text))
    return notes, hidden_skipped


def render_markdown(deck: Path, notes: list[tuple[int, str]]) -> str:
    lines = [f"# Speaker notes: {deck.name}", ""]
    for number, text in notes:
        lines.append(f"## Slide {number}")
        lines.append("")
        lines.append(text if text else "_(no notes)_")
        lines.append("")
    return "\n".join(lines)


def render_text(notes: list[tuple[int, str]]) -> str:
    lines = []
    for number, text in notes:
        lines.append(f"--- Slide {number} ---")
        lines.append(text if text else "(no notes)")
        lines.append("")
    return "\n".join(lines)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    parser.add_argument("deck", type=Path, help="Source .pptx deck")
    parser.add_argument(
        "--format",
        choices=("markdown", "text"),
        default="markdown",
        help="Output format (default: markdown)",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=None,
        help="Write to this file instead of stdout",
    )
    parser.add_argument(
        "--include-hidden",
        action="store_true",
        help=(
            "Keep hidden slides (default: skip them and renumber so the notes "
            "stay aligned with the extract_visuals.py PNGs)"
        ),
    )
    args = parser.parse_args()

    deck: Path = args.deck
    if not deck.exists():
        fail(f"Deck not found: {deck}")
    if deck.suffix.lower() != ".pptx":
        fail(
            f"Unsupported format {deck.suffix!r}: only .pptx is supported "
            "(export Keynote/Google Slides to .pptx first)"
        )

    notes, hidden_skipped = extract_notes(deck, include_hidden=args.include_hidden)
    out = render_markdown(deck, notes) if args.format == "markdown" else render_text(notes)

    # Always surface skipped hidden slides on stderr, so it never silently
    # pollutes piped stdout and the note<->PNG alignment stays auditable.
    if hidden_skipped:
        print(
            f"NOTE: skipped {hidden_skipped} hidden slide(s) to stay aligned with the "
            "extract_visuals.py PNGs (pass --include-hidden to keep them).",
            file=sys.stderr,
        )

    if args.output:
        args.output.parent.mkdir(parents=True, exist_ok=True)
        args.output.write_text(out + "\n")
        with_notes = sum(1 for _, text in notes if text)
        print(
            f"Wrote notes for {len(notes)} slides ({with_notes} with notes) to {args.output}",
            file=sys.stderr,
        )
    else:
        print(out)


if __name__ == "__main__":
    main()
